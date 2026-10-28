"""R-800: إنشاء ملفات حقيقية بصيغ متعددة من نص (Markdown مبسط) وجدول اختياري.

الصيغ: txt, md, html, json, csv, xlsx, docx, pdf (عربي مُشكَّل RTL بخط Noto), pptx, png.
المدخل: title، content (أسطر: «# عنوان»، «## فرعي»، «- بند»، فقرات)، table (صفوف؛ الأول رأس).
"""

from __future__ import annotations

import csv
import html
import io
import json
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

FONTS = Path(__file__).resolve().parent / "fonts"
FORMATS = {"txt": "text/plain", "md": "text/markdown", "html": "text/html",
           "json": "application/json", "csv": "text/csv",
           "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
           "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
           "pdf": "application/pdf",
           "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
           "png": "image/png"}
_AR = re.compile(r"[\u0600-\u06FF]")


def blocks(content: str) -> List[Tuple[str, str]]:
    out = []
    for raw in (content or "").splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if line.startswith("## "):
            out.append(("h2", line[3:].strip()))
        elif line.startswith("# "):
            out.append(("h1", line[2:].strip()))
        elif re.match(r"^\s*[-*•]\s+", line):
            out.append(("li", re.sub(r"^\s*[-*•]\s+", "", line)))
        else:
            out.append(("p", line.strip()))
    return out


def safe_filename(title: str, fmt: str) -> str:
    base = re.sub(r"[^\w\u0600-\u06FF\- ]+", "", title or "").strip().replace(" ", "_")[:80] or "file"
    return f"{base}.{fmt}"


def _rows(table: Optional[List[List[Any]]]) -> List[List[str]]:
    return [["" if c is None else str(c) for c in r] for r in (table or []) if isinstance(r, (list, tuple))]


def _html(title: str, content: str, table) -> str:
    rtl = bool(_AR.search((title or "") + (content or "")))
    parts = [f"<h1>{html.escape(title)}</h1>"] if title else []
    in_list = False
    for kind, text in blocks(content):
        if kind == "li" and not in_list:
            parts.append("<ul>")
            in_list = True
        if kind != "li" and in_list:
            parts.append("</ul>")
            in_list = False
        tag = {"h1": "h2", "h2": "h3", "li": "li", "p": "p"}[kind]
        parts.append(f"<{tag}>{html.escape(text)}</{tag}>")
    if in_list:
        parts.append("</ul>")
    rows = _rows(table)
    if rows:
        parts.append("<table border='1' cellpadding='6' style='border-collapse:collapse'>")
        for i, r in enumerate(rows):
            cell = "th" if i == 0 else "td"
            parts.append("<tr>" + "".join(f"<{cell}>{html.escape(c)}</{cell}>" for c in r) + "</tr>")
        parts.append("</table>")
    return (f"<!doctype html><html lang='{'ar' if rtl else 'en'}' dir='{'rtl' if rtl else 'ltr'}'>"
            f"<head><meta charset='utf-8'><title>{html.escape(title or '')}</title>"
            "<style>body{font-family:system-ui,'Noto Sans Arabic',sans-serif;max-width:860px;"
            "margin:2rem auto;padding:0 1rem;line-height:1.7}</style></head><body>"
            + "\n".join(parts) + "</body></html>")


def _md(title: str, content: str, table) -> str:
    out = [f"# {title}", ""] if title else []
    out.append(content or "")
    rows = _rows(table)
    if rows:
        out += ["", "| " + " | ".join(rows[0]) + " |", "|" + "---|" * len(rows[0])]
        out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out).strip() + "\n"


def _csv(rows) -> bytes:
    buf = io.StringIO()
    csv.writer(buf).writerows(rows)
    return ("\ufeff" + buf.getvalue()).encode("utf-8")  # BOM: Excel يقرأ العربية صحيحًا


def _xlsx(title: str, content: str, rows) -> bytes:
    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = (re.sub(r"[\[\]\*\?/\\:]", "", title or "Sheet") or "Sheet")[:31]
    if not rows:
        rows = [["المحتوى"]] + [[t] for _, t in blocks(content)]
    for r in rows:
        ws.append([_num(c) for c in r])
    for c in ws[1]:
        c.font = Font(bold=True)
    if _AR.search(json.dumps(rows, ensure_ascii=False)):
        ws.sheet_view.rightToLeft = True
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _num(v: str):
    try:
        return int(v) if re.fullmatch(r"-?\d+", v) else float(v) if re.fullmatch(r"-?\d+\.\d+", v) else v
    except ValueError:  # TOLERATED: يبقى نصًا
        return v


def _docx(title: str, content: str, rows) -> bytes:
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    d = docx.Document()
    rtl = bool(_AR.search((title or "") + (content or "")))

    def fix(p):
        if rtl:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            pPr = p._p.get_or_add_pPr()
            bidi = pPr.makeelement(qn("w:bidi"), {})
            pPr.append(bidi)
            for r in p.runs:
                r._r.get_or_add_rPr().append(r._r.makeelement(qn("w:rtl"), {}))
        return p
    if title:
        fix(d.add_heading(title, 0))
    for kind, text in blocks(content):
        if kind == "h1":
            fix(d.add_heading(text, 1))
        elif kind == "h2":
            fix(d.add_heading(text, 2))
        elif kind == "li":
            fix(d.add_paragraph(text, style="List Bullet"))
        else:
            fix(d.add_paragraph(text))
    if rows:
        t = d.add_table(rows=len(rows), cols=max(len(r) for r in rows))
        t.style = "Table Grid"
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                t.cell(i, j).text = c
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _pdf(title: str, content: str, rows) -> bytes:
    from fpdf import FPDF
    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(True, margin=15)
    pdf.add_font("NotoAr", fname=str(FONTS / "NotoSansArabic-Regular.ttf"))
    pdf.add_font("NotoLat", fname=str(FONTS / "NotoSans-Regular.ttf"))
    pdf.set_text_shaping(True)
    pdf.add_page()

    def para(text: str, size: float, gap: float = 2):
        ar = bool(_AR.search(text))
        pdf.set_font("NotoAr" if ar else "NotoLat", size=size)
        pdf.set_fallback_fonts(["NotoLat" if ar else "NotoAr"])
        pdf.multi_cell(0, size * 0.55, text, align="R" if ar else "L", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(gap)
    if title:
        para(title, 20, 4)
    for kind, text in blocks(content):
        para(("• " + text) if kind == "li" else text, {"h1": 16, "h2": 13.5}.get(kind, 11))
    if rows:
        pdf.set_font("NotoAr", size=10)
        pdf.set_fallback_fonts(["NotoLat"])
        from fpdf.fonts import FontFace
        with pdf.table(headings_style=FontFace(emphasis=None, fill_color=(226, 232, 240)),
                       text_align="CENTER") as t:
            for r in rows:
                row = t.row()
                for c in r:
                    row.cell(c)
    return bytes(pdf.output())


def _pptx(title: str, content: str, rows) -> bytes:
    from pptx import Presentation
    from pptx.util import Pt
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = title or "عرض"
    slides: List[Tuple[str, List[str]]] = []
    for kind, text in blocks(content):
        if kind in ("h1", "h2") or not slides:
            slides.append((text if kind in ("h1", "h2") else title or "", []))
            if kind in ("h1", "h2"):
                continue
        slides[-1][1].append(text)
    for head, items in slides:
        sl = prs.slides.add_slide(prs.slide_layouts[1])
        sl.shapes.title.text = head
        tf = sl.placeholders[1].text_frame
        tf.text = items[0] if items else ""
        for it in items[1:]:
            p = tf.add_paragraph()
            p.text = it
            p.font.size = Pt(18)
    if rows:
        sl = prs.slides.add_slide(prs.slide_layouts[5])
        sl.shapes.title.text = "جدول"
        n_c = max(len(r) for r in rows)
        shape = sl.shapes.add_table(len(rows), n_c, Pt(30), Pt(110), Pt(660), Pt(30 * len(rows)))
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                shape.table.cell(i, j).text = c
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _png(title: str, content: str) -> bytes:
    """بطاقة نصية حقيقية بالخط العربي المُشكَّل (Pillow + raqm إن توفر، وإلا تشكيل عبر fpdf غير متاح)."""
    from PIL import Image, ImageDraw, ImageFont, features
    W, H = 1200, 675
    im = Image.new("RGB", (W, H), (15, 23, 42))
    dr = ImageDraw.Draw(im)
    layout = ImageFont.Layout.RAQM if features.check("raqm") else ImageFont.Layout.BASIC
    ft = ImageFont.truetype(str(FONTS / "NotoSansArabic-Regular.ttf"), 54, layout_engine=layout)
    fb = ImageFont.truetype(str(FONTS / "NotoSansArabic-Regular.ttf"), 30, layout_engine=layout)
    mixed = {s: ImageFont.truetype(str(FONTS / "DejaVuSans.ttf"), s, layout_engine=layout)
             for s in (54, 30)}
    y = 70
    for text, font, color in [(title or "", ft, (248, 250, 252))] + [
            (t, fb, (203, 213, 225)) for _, t in blocks(content)[:10]]:
        if not text:
            continue
        if not re.fullmatch(r"[\u0600-\u06FF\s0-9.,:;!?()\-]*", text):  # R-801: Noto العربي بلا لاتيني ولا «—» — DejaVu يغطي الاثنين
            font = mixed[font.size]
        kw = {"direction": "rtl"} if layout == ImageFont.Layout.RAQM and _AR.search(text) else {}
        w = dr.textlength(text, font=font, **kw)
        x = W - 70 - w if _AR.search(text) else 70
        dr.text((x, y), text, font=font, fill=color, **kw)
        y += int(font.size * 1.6)
        if y > H - 60:
            break
    buf = io.BytesIO()
    im.save(buf, "PNG")
    return buf.getvalue()


def create(fmt: str, title: str = "", content: str = "",
           table: Optional[List[List[Any]]] = None, data: Any = None) -> Tuple[str, bytes, str]:
    fmt = (fmt or "").lower().lstrip(".")
    if fmt not in FORMATS:
        raise ValueError(f"صيغة غير مدعومة: {fmt} — المتاح: {', '.join(sorted(FORMATS))}")
    rows = _rows(table)
    if fmt == "txt":
        blob = ((title + "\n\n") if title else "").encode() + (content or "").encode("utf-8")
    elif fmt == "md":
        blob = _md(title, content, rows).encode("utf-8")
    elif fmt == "html":
        blob = _html(title, content, rows).encode("utf-8")
    elif fmt == "json":
        blob = json.dumps(data if data is not None else {"title": title, "content": content,
                                                          "table": rows}, ensure_ascii=False,
                          indent=2).encode("utf-8")
    elif fmt == "csv":
        blob = _csv(rows or [[t] for _, t in blocks(content)])
    elif fmt == "xlsx":
        blob = _xlsx(title, content, rows)
    elif fmt == "docx":
        blob = _docx(title, content, rows)
    elif fmt == "pdf":
        blob = _pdf(title, content, rows)
    elif fmt == "pptx":
        blob = _pptx(title, content, rows)
    else:
        blob = _png(title, content)
    return safe_filename(title, fmt), blob, FORMATS[fmt]
