"""R-800: قراءة الملفات الحقيقية واستخراج محتواها وبنيتها.

extract(name, data) ⇒ {"kind", "supported", "text", "meta", "structure", "error"}
- النص يُقص عند MAX_TEXT حرف ويُعلن truncated=true.
- الصيغ غير المدعومة تُعلن supported=false بالسبب — لا محتوى مختلق.
"""

from __future__ import annotations

import csv
import io
import json
import zipfile
from pathlib import PurePosixPath
from typing import Any, Dict, List

MAX_TEXT = 200_000

TEXT_EXT = {".txt", ".md", ".markdown", ".rst", ".log", ".ini", ".cfg", ".conf", ".toml", ".yaml",
            ".yml", ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".kt", ".go", ".rs", ".c", ".h",
            ".cpp", ".hpp", ".cs", ".rb", ".php", ".sh", ".bash", ".sql", ".css", ".scss", ".swift",
            ".r", ".lua", ".pl", ".dart", ".vue", ".svelte", ".env.example", ".gitignore",
            ".dockerfile", ".tf", ".gradle", ".properties", ".srt", ".vtt", ".tex"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".bmp", ".tiff", ".tif", ".ico"}
AUDIO_EXT = {".mp3", ".wav", ".ogg", ".m4a", ".flac", ".webm", ".aac"}
VIDEO_EXT = {".mp4", ".mov", ".mkv", ".avi"}
ARCHIVE_EXT = {".zip", ".tar", ".tgz", ".gz", ".bz2", ".xz", ".7z", ".tbz2", ".txz"}


def ext_of(name: str) -> str:
    n = PurePosixPath(name.lower()).name
    for multi in (".tar.gz", ".tar.bz2", ".tar.xz", ".env.example"):
        if n.endswith(multi):
            return multi
    if n == "dockerfile":
        return ".dockerfile"
    return PurePosixPath(n).suffix


def kind_of(name: str, data: bytes = b"") -> str:
    e = ext_of(name)
    if e in (".tar.gz", ".tar.bz2", ".tar.xz") or e in ARCHIVE_EXT:
        return "archive"
    if e == ".pdf" or data[:5] == b"%PDF-":
        return "pdf"
    if e == ".docx":
        return "docx"
    if e in (".xlsx", ".xlsm"):
        return "xlsx"
    if e == ".pptx":
        return "pptx"
    if e in (".csv", ".tsv"):
        return "csv"
    if e == ".json":
        return "json"
    if e in (".html", ".htm", ".xml", ".svg"):
        return "markup"
    if e in IMAGE_EXT:
        return "image"
    if e in AUDIO_EXT:
        return "audio"
    if e in VIDEO_EXT:
        return "video"
    if e in TEXT_EXT:
        return "text"
    if data[:4] == b"PK\x03\x04":
        return "archive"
    try:
        data[:4096].decode("utf-8")
        return "text"
    except UnicodeDecodeError:
        return "binary"


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "cp1256", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:  # TOLERATED: نجرب الترميز التالي؛ latin-1 لا يفشل
            continue
    return data.decode("utf-8", "replace")


def _cap(text: str) -> Dict[str, Any]:
    return {"text": text[:MAX_TEXT], "truncated": len(text) > MAX_TEXT, "chars": len(text)}


def _text(data: bytes) -> Dict[str, Any]:
    t = _decode(data)
    lines = t.splitlines()
    return {**_cap(t), "structure": {"lines": len(lines), "words": len(t.split()),
                                     "blank_lines": sum(1 for l in lines if not l.strip())}}


def _csv(name: str, data: bytes) -> Dict[str, Any]:
    t = _decode(data)
    delim = "\t" if ext_of(name) == ".tsv" else (csv.Sniffer().sniff(t[:4096]).delimiter
                                                  if t.strip() else ",")
    rows = list(csv.reader(io.StringIO(t), delimiter=delim))
    header = rows[0] if rows else []
    body = rows[1:]
    cols = []
    for i, h in enumerate(header[:100]):
        vals = [r[i] for r in body if i < len(r) and r[i] != ""]
        nums = []
        for v in vals:
            try:
                nums.append(float(v.replace(",", "")))
            except ValueError:  # TOLERATED: قيمة غير رقمية — العمود نصي جزئيًا
                pass
        c = {"name": h, "non_empty": len(vals), "unique": len(set(vals))}
        if nums and len(nums) >= 0.8 * len(vals):
            c.update(numeric=True, min=min(nums), max=max(nums), mean=round(sum(nums) / len(nums), 4))
        cols.append(c)
    return {**_cap(t), "structure": {"rows": len(body), "columns": len(header), "delimiter": delim,
                                     "header": header[:100], "column_stats": cols,
                                     "sample": body[:5]}}


def _json(data: bytes) -> Dict[str, Any]:
    t = _decode(data)
    obj = json.loads(t)

    def shape(o, depth=0):
        if depth > 3:
            return type(o).__name__
        if isinstance(o, dict):
            return {k: shape(v, depth + 1) for k, v in list(o.items())[:30]}
        if isinstance(o, list):
            return [f"list[{len(o)}]", shape(o[0], depth + 1) if o else None]
        return type(o).__name__
    return {**_cap(t), "structure": {"top_type": type(obj).__name__, "shape": shape(obj)}}


def _markup(data: bytes) -> Dict[str, Any]:
    import re
    t = _decode(data)
    title = re.search(r"<title[^>]*>(.*?)</title>", t, re.S | re.I)
    plain = re.sub(r"<script.*?</script>|<style.*?</style>", " ", t, flags=re.S | re.I)
    plain = re.sub(r"<[^>]+>", " ", plain)
    plain = re.sub(r"\s+", " ", plain).strip()
    links = re.findall(r'href=["\']([^"\']+)', t)[:50]
    return {**_cap(plain), "structure": {"title": title.group(1).strip() if title else None,
                                         "links": links, "raw_chars": len(t)}}


_BIDI_MARKS = dict.fromkeys(map(ord, "\u200e\u200f\u202a\u202b\u202c\u202d\u202e\u2066\u2067\u2068\u2069"))


def _poppler_pages(data: bytes, n: int) -> List[str]:
    """R-800: pdftotext (poppler) يحفظ ترتيب النص العربي المختلط بالأرقام؛ pypdf يُسقط منه
    مقاطع (مثبت باختبار). يُستعمل إن وُجد في الحاوية، وإلا يبقى pypdf وحده."""
    import shutil
    import subprocess
    if not shutil.which("pdftotext"):
        return []
    try:
        out = subprocess.run(["pdftotext", "-enc", "UTF-8", "-l", str(min(n, 500)), "-", "-"], input=data,
                             capture_output=True, timeout=60).stdout.decode("utf-8", "replace")
    except (OSError, subprocess.SubprocessError):
        return []
    pages = [pg.translate(_BIDI_MARKS).strip() for pg in out.split("\f")]
    return pages[:n] if pages and any(pages) else []


def _pdf(data: bytes) -> Dict[str, Any]:
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    pages = []
    for i, p in enumerate(r.pages[:500]):
        try:
            pages.append(p.extract_text() or "")
        except Exception as exc:  # TOLERATED: صفحة تالفة تُعلن نصًا فارغًا بسببها
            pages.append(f"[تعذر استخراج الصفحة {i + 1}: {type(exc).__name__}]")
    pop = _poppler_pages(data, len(pages))
    engine = "pypdf"
    if pop and sum(len(t) for t in pop) >= sum(len(t) for t in pages) * 0.9:
        pages = pop + pages[len(pop):]
        engine = "poppler"
    text = "\n\n".join(f"--- صفحة {i + 1} ---\n{t}" for i, t in enumerate(pages))
    meta = {k.strip("/"): str(v) for k, v in (r.metadata or {}).items()} if r.metadata else {}
    empty = sum(1 for t in pages if not t.strip())
    return {**_cap(text), "meta": meta,
            "structure": {"pages": len(r.pages), "pages_without_text": empty,
                          "encrypted": bool(r.is_encrypted), "engine": engine,
                          "note": "صفحات بلا نص غالبًا صور ممسوحة — تحتاج OCR/رؤية" if empty else None}}


def _docx(data: bytes) -> Dict[str, Any]:
    import docx
    d = docx.Document(io.BytesIO(data))
    parts, headings = [], []
    for p in d.paragraphs:
        if p.text.strip():
            parts.append(p.text)
            if p.style is not None and str(p.style.name).lower().startswith(("heading", "title")):
                headings.append(p.text[:120])
    tables = []
    for t in d.tables[:50]:
        rows = [[c.text for c in r.cells] for r in t.rows]
        tables.append({"rows": len(rows), "cols": len(rows[0]) if rows else 0, "sample": rows[:5]})
        parts.append("\n".join(" | ".join(r) for r in rows))
    cp = d.core_properties
    return {**_cap("\n".join(parts)), "meta": {"title": cp.title, "author": cp.author},
            "structure": {"paragraphs": len(d.paragraphs), "headings": headings[:100],
                          "tables": tables, "images": len(d.inline_shapes)}}


def _xlsx(data: bytes) -> Dict[str, Any]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets, parts = [], []
    for ws in wb.worksheets[:30]:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 2000:
                break
            rows.append(["" if v is None else str(v) for v in row])
        sheets.append({"name": ws.title, "rows_read": len(rows), "max_row": ws.max_row,
                       "max_column": ws.max_column, "sample": rows[:5]})
        parts.append(f"### {ws.title}\n" + "\n".join("\t".join(r) for r in rows))
    return {**_cap("\n\n".join(parts)), "structure": {"sheets": sheets}}


def _pptx(data: bytes) -> Dict[str, Any]:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    slides, parts = [], []
    for i, s in enumerate(prs.slides):
        texts = [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()]
        notes = s.notes_slide.notes_text_frame.text if s.has_notes_slide else ""
        slides.append({"n": i + 1, "title": texts[0][:120] if texts else None, "shapes": len(s.shapes)})
        parts.append(f"--- شريحة {i + 1} ---\n" + "\n".join(texts) + (f"\n[ملاحظات] {notes}" if notes else ""))
    return {**_cap("\n\n".join(parts)), "structure": {"slides": len(prs.slides), "outline": slides}}


def _image(data: bytes) -> Dict[str, Any]:
    from PIL import Image, ExifTags
    im = Image.open(io.BytesIO(data))
    exif = {}
    try:
        for k, v in (im.getexif() or {}).items():
            exif[ExifTags.TAGS.get(k, str(k))] = str(v)[:120]
    except Exception:  # TOLERATED: EXIF تالف — يُترك فارغًا
        exif = {}
    small = im.convert("RGB").resize((32, 32))
    px = list(small.getdata())
    avg = tuple(round(sum(c[i] for c in px) / len(px)) for i in range(3))
    return {"text": "", "truncated": False, "chars": 0,
            "meta": {"format": im.format, "mode": im.mode, "exif": exif},
            "structure": {"width": im.width, "height": im.height,
                          "frames": getattr(im, "n_frames", 1), "average_rgb": avg,
                          "note": "وصف المحتوى المرئي يتطلب نموذج رؤية (analyze)"}}


def _archive(name: str, data: bytes) -> Dict[str, Any]:
    from . import archives
    listing = archives.list_members(name, data)
    text = "\n".join(f"{m['name']}\t{m['size']}" for m in listing["members"])
    return {**_cap(text), "structure": listing}


def extract(name: str, data: bytes) -> Dict[str, Any]:
    kind = kind_of(name, data)
    base: Dict[str, Any] = {"kind": kind, "supported": True, "text": "", "truncated": False,
                            "chars": 0, "meta": {}, "structure": {}, "error": None,
                            "bytes": len(data)}
    handlers = {"text": lambda: _text(data), "csv": lambda: _csv(name, data),
                "json": lambda: _json(data), "markup": lambda: _markup(data),
                "pdf": lambda: _pdf(data), "docx": lambda: _docx(data), "xlsx": lambda: _xlsx(data),
                "pptx": lambda: _pptx(data), "image": lambda: _image(data),
                "archive": lambda: _archive(name, data)}
    h = handlers.get(kind)
    if h is None:
        base["supported"] = False
        base["error"] = {"audio": "الصوت يُفرغ نصيًا عبر محرك STT (analyze) إن كان مفعّلًا",
                         "video": "الفيديو: تُستخرج البيانات الوصفية عبر ffprobe (analyze)",
                         "binary": "ملف ثنائي غير معروف — يُخزن ويُنزّل كما هو"}.get(kind)
        return base
    try:
        base.update(h())
    except Exception as exc:  # TOLERATED: ملف تالف/محمي — الخطأ الحرفي يُعاد للمستخدم
        base["supported"] = False
        base["error"] = f"{type(exc).__name__}: {str(exc)[:300]}"
    return base
